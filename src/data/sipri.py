"""SIPRI Military Expenditure Database adapter.

The official workbook (``SIPRI-Milex-data-1949-2025_v1.2.xlsx``) is a
presentation spreadsheet: one sheet per indicator, countries grouped under
region/subregion header rows, missing values written as "..." and "xxx", and the
estimate/uncertainty status of each figure encoded only as *font colour*
(blue = SIPRI estimate, red = highly uncertain). We read cell values *and* fonts
with openpyxl so that those flags survive into the analytical tables.
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd

from src.data import countries
from src.data.schema import MILEX_COLUMNS, MILEX_INDICATORS, conform
from src.data.validation import QualityReport

EUROPE_SUBREGIONS = {"Central Europe", "Eastern Europe", "Western Europe"}
OTHER_REGIONS = {"Africa", "Americas", "Asia & Oceania", "Middle East"}
AGGREGATES = {"European Union": "EUU"}

# openpyxl indexed colours used by SIPRI: 12 = blue, 10 = red.
_BLUE, _RED = {12}, {10}


def _flag_from_font(cell) -> str:
    color = cell.font.color if cell.font is not None else None
    if color is None:
        return "reported"
    if color.type == "indexed":
        if color.indexed in _BLUE:
            return "estimate"
        if color.indexed in _RED:
            return "highly_uncertain"
    if color.type == "rgb" and isinstance(color.rgb, str):
        rgb = color.rgb.upper()[-6:]
        if rgb == "0000FF":
            return "estimate"
        if rgb == "FF0000":
            return "highly_uncertain"
    return "reported"


def _find_sheet(wb, wanted: str) -> tuple[str | None, int | None]:
    """Return the actual sheet name and, for constant prices, the base year."""
    if wanted.startswith("Constant"):
        for name in wb.sheetnames:
            m = re.match(r"Constant \((\d{4})\) US\$", name)
            if m:
                return name, int(m.group(1))
        return None, None
    return (wanted if wanted in wb.sheetnames else None), None


def _header(ws) -> tuple[int, dict[int, int]]:
    """Locate the 'Country' header row and map column index -> year."""
    for row in ws.iter_rows(min_row=1, max_row=15):
        if row[0].value and str(row[0].value).strip() == "Country":
            years = {}
            for cell in row[1:]:
                try:
                    y = int(float(str(cell.value).strip()))
                except (TypeError, ValueError):
                    continue
                if 1900 <= y <= 2100:
                    years[cell.column - 1] = y
            return row[0].row, years
    raise ValueError(f"No 'Country' header row in sheet {ws.title!r}")


def _parse_sheet(ws, indicator: str, scale: float, report: QualityReport) -> list[dict]:
    header_row, years = _header(ws)
    meta = MILEX_INDICATORS[indicator]
    rows, region, sub, unmapped = [], None, None, set()
    for row in ws.iter_rows(min_row=header_row + 1):
        name = row[0].value
        if name is None:
            continue
        name = str(name).strip()
        values = [row[c] for c in years if c < len(row)]
        is_header = all(v.value in (None, "") for v in values)
        if is_header:
            if name == "Europe":
                region = "Europe"
            elif name in OTHER_REGIONS:
                region = name
            if name in EUROPE_SUBREGIONS:
                sub = name
            continue
        if region != "Europe":
            continue
        iso = AGGREGATES.get(name) or countries.to_iso3(name)
        if iso is None:
            unmapped.add(name)
            continue
        entity = (
            "aggregate"
            if name in AGGREGATES
            else "historical" if iso in countries.HISTORICAL else "country"
        )
        for col, year in years.items():
            if col >= len(row):
                continue
            cell = row[col]
            v = cell.value
            if isinstance(v, (int, float)):
                value, flag = float(v) * scale, _flag_from_font(cell)
            else:
                s = str(v).replace(" ", "") if v is not None else ""
                value = None
                flag = "not_independent" if s.lower() == "xxx" else "unavailable"
            rows.append(
                {
                    "country": name if entity == "aggregate" else countries.display_name(iso),
                    "iso3": iso,
                    "subregion": None if entity == "aggregate" else sub,
                    "entity_type": entity,
                    "year": year,
                    "indicator": indicator,
                    "unit": meta["unit"],
                    "price_basis": meta["price_basis"],
                    "value": value,
                    "value_flag": flag,
                }
            )
    report.add(
        f"country_harmonisation[{indicator}]",
        "warn" if unmapped else "pass",
        len(unmapped),
        len(unmapped) + len({r["iso3"] for r in rows}),
        f"Unmapped European entries: {sorted(unmapped)}" if unmapped else "All European rows mapped.",
        "unmapped rows dropped" if unmapped else "none",
    )
    return rows


def load_sipri(path: Path) -> tuple[pd.DataFrame, QualityReport]:
    report = QualityReport(source="sipri_milex", input_file=Path(path).name)
    wb = openpyxl.load_workbook(path, data_only=True)
    all_rows = []
    for key, meta in MILEX_INDICATORS.items():
        sheet, base_year = _find_sheet(wb, meta["sheet"])
        if sheet is None:
            report.add(f"sheet_present[{key}]", "fail", 1, 1, f"Sheet {meta['sheet']!r} not found.")
            continue
        report.add(f"sheet_present[{key}]", "pass", 0, 1, f"Read sheet {sheet!r}.")
        rows = _parse_sheet(wb[sheet], key, float(meta["scale"]), report)
        if base_year:
            for r in rows:
                r["price_basis"] = f"constant {base_year} prices and exchange rates"
        all_rows += rows
    df = pd.DataFrame(all_rows)
    df["source"] = "sipri_milex"
    df["data_status"] = "real"
    report.rows_in = len(df)

    dup = df.duplicated(subset=["iso3", "year", "indicator"])
    report.add(
        "duplicate_country_year_indicator",
        "fail" if dup.any() else "pass",
        int(dup.sum()),
        len(df),
        "Each (country, year, indicator) must appear once.",
    )
    df = df.loc[~dup]
    neg = df["value"].notna() & (df["value"] < 0)
    report.add("values_non_negative", "fail" if neg.any() else "pass", int(neg.sum()), len(df),
               "Expenditure and shares cannot be negative.")  # fmt: skip
    share = df["indicator"].isin(["share_gdp_pct", "share_govt_pct"]) & (df["value"] > 100)
    report.add("shares_at_most_100pct", "warn" if share.any() else "pass", int(share.sum()),
               len(df), "Shares above 100% indicate a unit error.")  # fmt: skip
    for flag in ["unavailable", "not_independent", "estimate", "highly_uncertain"]:
        k = int((df["value_flag"] == flag).sum())
        report.add(f"value_flag_{flag}", "info", k, len(df), f"{k} cells flagged '{flag}'.")
    report.transformations += [
        "Selected the Europe region (Central, Eastern, Western Europe) and the EU aggregate.",
        "Wide year columns reshaped to long format (one row per country-year-indicator).",
        "'...' -> missing (value_flag=unavailable); 'xxx' -> missing (not_independent). Never 0.",
        "Blue font -> value_flag=estimate; red font -> value_flag=highly_uncertain.",
        "Share of GDP and share of government spending converted from fractions to percent (x100).",
    ]
    out = conform(df, MILEX_COLUMNS)
    report.rows_out = len(out)
    have = out.dropna(subset=["value"])
    report.coverage = {
        "year_min": int(have["year"].min()),
        "year_max": int(have["year"].max()),
        "n_entities": int(out["iso3"].nunique()),
        "indicators": sorted(out["indicator"].unique().tolist()),
    }
    return out, report


def load_milex_csv(path: Path, source: str = "synthetic_milex") -> tuple[pd.DataFrame, QualityReport]:
    """Load a long-format expenditure CSV that already follows MILEX_COLUMNS."""
    report = QualityReport(source=source, input_file=Path(path).name)
    df = pd.read_csv(path)
    report.rows_in = len(df)
    missing = [c for c in MILEX_COLUMNS if c not in df.columns]
    report.add("required_columns", "fail" if missing else "pass", len(missing), len(MILEX_COLUMNS),
               f"Missing: {missing}" if missing else "All columns present.")  # fmt: skip
    if missing:
        raise ValueError(f"{path}: missing columns {missing}")
    if source == "synthetic_milex" and not (df["data_status"].str.lower() == "synthetic").all():
        raise ValueError("Synthetic expenditure file must label every row data_status=synthetic")
    bad_ind = ~df["indicator"].isin(list(MILEX_INDICATORS))
    report.add("known_indicators", "fail" if bad_ind.any() else "pass", int(bad_ind.sum()), len(df),
               "Indicators must be in the catalogue.")  # fmt: skip
    df = df.loc[~bad_ind]
    out = conform(df, MILEX_COLUMNS)
    report.rows_out = len(out)
    report.coverage = {"year_min": int(out["year"].min()), "year_max": int(out["year"].max()),
                       "n_entities": int(out["iso3"].nunique())}  # fmt: skip
    return out, report
