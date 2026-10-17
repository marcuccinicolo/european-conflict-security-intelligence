"""Shared fixtures: tiny hand-built inputs whose correct outputs are known."""

from __future__ import annotations

import openpyxl
import pandas as pd
import pytest
from openpyxl.styles import Font

from src.data.schema import EVENT_COLUMNS, conform


def make_events(rows: list[dict], source: str = "ucdp_ged") -> pd.DataFrame:
    base = {
        "source": source, "data_status": "real", "country": "Ukraine", "iso3": "UKR",
        "latitude": 50.0, "longitude": 30.0, "coord_status": "valid", "coord_check": "inside",
        "location_precise": True, "date_uncertain": False, "event_type": "State-based conflict",
        "fatalities_kind": "estimate",
    }  # fmt: skip
    out = []
    for i, r in enumerate(rows):
        d = {**base, "event_id": f"e{i}", **r}
        d["event_date"] = pd.Timestamp(d["event_date"])
        out.append(d)
    df = conform(pd.DataFrame(out), EVENT_COLUMNS)
    df["year"] = df["event_date"].dt.year
    return df


@pytest.fixture
def acled_csv(tmp_path):
    """ACLED-format CSV with one of each defect."""
    rows = [
        # id,        date,            type,     sub,           country,  lat,   lon,   fat
        ("A1", "2024-01-05", "Battles", "Armed clash", "Ukraine", "48.0", "37.8", "3"),
        ("A2", "06 February 2024", "Protests", "Peaceful protest", "France", "48.85", "2.35", "0"),
        ("A3", "07/03/2024", "Riots", "Mob violence", "Germany", "52.5", "13.4", ""),
        ("A4", "2024-13-40", "Battles", "Armed clash", "Ukraine", "48.0", "37.8", "1"),
        ("A5", "2024-04-01", "Battles", "Armed clash", "Ukraine", "", "", "-2"),
        ("A6", "2024-04-02", "Battles", "Armed clash", "Ukraine", "95", "37.8", "2"),
        ("A7", "2024-04-03", "Battles", "Armed clash", "Narnia", "48.0", "37.8", "2"),
        ("A1", "2024-01-05", "Battles", "Armed clash", "Ukraine", "48.0", "37.8", "3"),
    ]  # fmt: skip
    df = pd.DataFrame(rows, columns=["event_id_cnty", "event_date", "event_type", "sub_event_type",
                                     "country", "latitude", "longitude", "fatalities"])  # fmt: skip
    df["admin1"] = "Area"
    df["geo_precision"] = "1"
    df["time_precision"] = "1"
    p = tmp_path / "acled.csv"
    df.to_csv(p, index=False)
    return p


@pytest.fixture
def ucdp_csv(tmp_path):
    cols = ["id", "year", "type_of_violence", "conflict_name", "where_prec", "adm_1", "latitude",
            "longitude", "country", "region", "date_prec", "date_start", "date_end", "best",
            "low", "high", "number_of_sources"]  # fmt: skip
    rows = [
        [1, 2022, 1, "Russia - Ukraine", 1, "Kharkiv", 49.99, 36.23, "Ukraine", "Europe", 1,
         "2022-03-01 00:00:00.000", "2022-03-01 00:00:00.000", 5, 4, 8, 2],
        [2, 2022, 3, "Russia - civilians", 4, "", 50.6, 36.6, "Russia (Soviet Union)", "Europe", 3,
         "2022-03-10 00:00:00.000", "2022-03-16 00:00:00.000", 2, 2, 2, -1],
        [3, 2022, 1, "Syria", 1, "Aleppo", 36.2, 37.1, "Syria", "Middle East", 1,
         "2022-03-01 00:00:00.000", "2022-03-01 00:00:00.000", 1, 1, 1, 1],
        [4, 2023, 2, "X - Y", 2, "Kosovo", 42.66, 21.16, "Serbia (Yugoslavia)", "Europe", 1,
         "2023-05-01 00:00:00.000", "2023-05-01 00:00:00.000", 3, 5, 3, 1],
    ]  # fmt: skip
    p = tmp_path / "GEDEvent_test.csv"
    pd.DataFrame(rows, columns=cols).to_csv(p, index=False)
    return p


@pytest.fixture
def sipri_xlsx(tmp_path):
    """A miniature SIPRI workbook reproducing the real layout quirks."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    blue, red = Font(color="FF0000FF"), Font(color="FFFF0000")

    def sheet(name, header_row, values, notes_col=True):
        ws = wb.create_sheet(name)
        ws.cell(row=1, column=1, value=f"{name} © SIPRI")
        ws.cell(row=header_row, column=1, value="Country")
        start = 3 if notes_col else 2
        if notes_col:
            ws.cell(row=header_row, column=2, value="Notes")
        for j, y in enumerate([2022, 2023, 2024]):
            ws.cell(row=header_row, column=start + j, value=y)
        r = header_row + 2
        for name_, vals, fonts in values:
            ws.cell(row=r, column=1, value=name_)
            for j, v in enumerate(vals):
                c = ws.cell(row=r, column=start + j, value=v)
                if fonts and fonts[j]:
                    c.font = fonts[j]
            r += 1

    rows = [
        ("Middle East", [None, None, None], None),
        ("Israel", [100.0, 110.0, 120.0], None),
        ("Europe", [None, None, None], None),
        ("Western Europe", [None, None, None], None),
        ("Germany", [50000.0, 55000.0, 60000.0], [None, None, blue]),
        ("France", ["...", 40000.0, 42000.0], [None, None, red]),
        ("Eastern Europe", [None, None, None], None),
        ("Ukraine", [40000.0, 60000.0, 65000.0], [blue, blue, blue]),
        ("USSR", ["xxx", "xxx", "xxx"], None),
        ("European Union", [300000.0, 320000.0, 340000.0], None),
    ]
    sheet("Constant (2024) US$", 6, rows, notes_col=True)
    sheet("Current US$", 6, rows)
    gdp = [(n, [v / 1e6 if isinstance(v, float) else v for v in vals], f) for n, vals, f in rows]
    sheet("Share of GDP", 6, gdp)
    p = tmp_path / "SIPRI-Milex-data-test.xlsx"
    wb.save(p)
    return p
